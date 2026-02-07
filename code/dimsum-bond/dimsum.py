from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import date


def build_dimsum_pnl_attribution_xlsx(out_path: str) -> None:
    """
    Creates an Excel workbook for Dim Sum (CNH) bond P&L attribution using:
      - Base curve: China sovereign CNH curve nodes 1/2/3/5/7/10
      - Spread: Z-spread (bond and class bucket)
      - Buckets: Rating + Sector + Tenor (rounded YEARFRAC(issue, maturity))
    All derived fields are computed by Excel formulas (openpyxl writes formulas; Excel evaluates).
    """

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    # -----------------------------
    # Styles
    # -----------------------------
    thin = Side(style="thin", color="D9D9D9")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=11)
    header_font = Font(bold=True, color="FFFFFF")

    header_fill = PatternFill("solid", fgColor="4F81BD")  # blue header
    input_fill = PatternFill("solid", fgColor="FFF2CC")  # light yellow inputs

    wrap_left = Alignment(wrap_text=True, vertical="top", horizontal="left")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    # -----------------------------
    # Sheets
    # -----------------------------
    ws_readme = wb.create_sheet("README")
    ws_inputs = wb.create_sheet("Inputs")
    ws_bucket = wb.create_sheet("Bucket_Data")
    ws_attr = wb.create_sheet("Attribution")

    # -----------------------------
    # README
    # -----------------------------
    ws_readme["A1"] = "Dim Sum Bond P&L Attribution (Campisi-style, Z-spread)"
    ws_readme["A1"].font = title_font

    ws_readme["A3"] = "Purpose"
    ws_readme["A3"].font = section_font
    ws_readme["A4"] = (
        "This workbook attributes realized P&L for a CNH (Dim Sum) bond into: "
        "Income (accrual+cash coupon), Base Curve (China sovereign CNH curve nodes 1/2/3/5/7/10), "
        "Class Spread, Selection (bond vs class), and Residual. It uses Z-spread over the base curve."
    )
    ws_readme["A4"].alignment = wrap_left

    ws_readme["A6"] = "How to use"
    ws_readme["A6"].font = section_font
    ws_readme["A7"] = (
        "1) Fill the yellow 'Inputs' cells with your actual market data (clean prices, curve node yields, "
        "Z-spreads, key-rate DV01s, spread DV01). "
        "2) 'Bucket_Data' holds class Z-spreads by Rating+Sector+Tenor bucket; update as needed. "
        "3) 'Attribution' computes P&L in price points and currency."
    )
    ws_readme["A7"].alignment = wrap_left

    ws_readme["A9"] = "Notes"
    ws_readme["A9"].font = section_font
    ws_readme["A10"] = (
        "• This file includes illustrative market inputs. Replace with Bloomberg/LSEG/your pricer outputs.\n"
        "• Class Z-spreads are bucket averages by Rating|Sector|TenorY. TenorY is rounded YEARFRAC(issue, maturity)."
    )
    ws_readme["A10"].alignment = wrap_left
    ws_readme.column_dimensions["A"].width = 120

    # -----------------------------
    # Bucket_Data
    # -----------------------------
    ws_bucket["A1"] = "Class / Bucket Z-Spreads (Rating + Sector + Tenor bucket)"
    ws_bucket["A1"].font = title_font

    bucket_headers = ["Bucket key (Rating|Sector|TenorY)", "Class Z t0 (bp)", "Class Z t1 (bp)", "Notes"]
    for col, h in enumerate(bucket_headers, start=1):
        c = ws_bucket.cell(row=2, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border_thin

    # Example bucket row (replace with your own bucket averages)
    ws_bucket["A3"] = "A-|Internet Services|5"
    ws_bucket["B3"] = 130
    ws_bucket["C3"] = 150
    ws_bucket["D3"] = "Illustrative bucket z-spreads. Replace with your bucket averages (rating+sector+tenor)."

    for col in range(1, 5):
        ws_bucket.cell(row=3, column=col).border = border_thin
        ws_bucket.cell(row=3, column=col).alignment = left

    ws_bucket["B3"].fill = input_fill
    ws_bucket["C3"].fill = input_fill
    ws_bucket["B3"].number_format = "0"
    ws_bucket["C3"].number_format = "0"

    ws_bucket.column_dimensions["A"].width = 32
    ws_bucket.column_dimensions["B"].width = 16
    ws_bucket.column_dimensions["C"].width = 16
    ws_bucket.column_dimensions["D"].width = 70
    ws_bucket.freeze_panes = "A3"

    # -----------------------------
    # Inputs
    # -----------------------------
    ws_inputs["A1"] = "INPUTS (yellow cells are user inputs)"
    ws_inputs["A1"].font = title_font

    ws_inputs["A3"] = "Bond Static (HK0001241626)"
    ws_inputs["A3"].font = section_font

    # Static inputs for sample bond HK0001241626 (replace if needed)
    static_rows = [
        ("ISIN", "HK0001241626"),
        ("Issuer", "Kuaishou Technology"),
        ("Bond name (short)", "KUAISHOU TEC 26/31"),
        ("Currency (treat as CNH for P&L)", "CNH"),
        ("Issue date", date(2026, 1, 22)),
        ("Maturity date", date(2031, 1, 22)),
        ("Coupon rate (annual)", 0.0245),
        ("Coupon frequency (payments/year)", 2),
        ("Day count", "ACT/365"),
        ("Issue size (principal)", 3_500_000_000),
        ("Rating (for bucket key)", "A-"),
        ("Sector (for bucket key)", "Internet Services"),
        ("Tenor bucket (years)", None),  # formula below
    ]

    start_row = 4
    for i, (k, v) in enumerate(static_rows):
        r = start_row + i
        ws_inputs[f"A{r}"] = k
        ws_inputs[f"B{r}"] = v
        ws_inputs[f"A{r}"].border = border_thin
        ws_inputs[f"B{r}"].border = border_thin
        ws_inputs[f"A{r}"].alignment = left
        ws_inputs[f"B{r}"].alignment = left

    # Tenor bucket formula: rounded YEARFRAC(issue,maturity)
    ws_inputs["B16"].value = "=ROUND(YEARFRAC(B8,B9,1),0)"
    ws_inputs["B8"].number_format = "yyyy-mm-dd"
    ws_inputs["B9"].number_format = "yyyy-mm-dd"
    ws_inputs["B10"].number_format = "0.0000%"
    ws_inputs["B13"].number_format = "#,##0"
    ws_inputs["B16"].number_format = "0"

    ws_inputs["A18"] = "Valuation Period & Position"
    ws_inputs["A18"].font = section_font

    val_rows = [
        ("Valuation date t0", date(2026, 1, 26)),
        ("Valuation date t1", date(2026, 1, 30)),
        ("Last coupon date (for accrual)", date(2026, 1, 22)),
        ("Next coupon date", date(2026, 7, 22)),
        ("Notional (face amount, CNH)", 100_000_000),
        ("Funding cost (CNH, period)", 0),
    ]
    start_row = 19
    for i, (k, v) in enumerate(val_rows):
        r = start_row + i
        ws_inputs[f"A{r}"] = k
        ws_inputs[f"B{r}"] = v
        ws_inputs[f"A{r}"].border = border_thin
        ws_inputs[f"B{r}"].border = border_thin

        if "date" in k.lower() or "coupon" in k.lower():
            ws_inputs[f"B{r}"].number_format = "yyyy-mm-dd"

        if "Notional" in k:
            ws_inputs[f"B{r}"].number_format = "#,##0"
            ws_inputs[f"B{r}"].fill = input_fill

        if "Funding cost" in k:
            ws_inputs[f"B{r}"].number_format = "#,##0"
            ws_inputs[f"B{r}"].fill = input_fill

    ws_inputs["A26"] = "Market Quotes (illustrative - replace with your marks)"
    ws_inputs["A26"].font = section_font

    quote_rows = [
        ("Clean price at t0 (per 100)", 100.10),
        ("Clean price at t1 (per 100)", 99.80),
        ("Coupon cash paid during period (per 100)", 0),
        ("Accrued interest at t0 (per 100)", None),
        ("Accrued interest at t1 (per 100)", None),
        ("Dirty price at t0 (per 100)", None),
        ("Dirty price at t1 (per 100)", None),
    ]
    start_row = 27
    for i, (k, v) in enumerate(quote_rows):
        r = start_row + i
        ws_inputs[f"A{r}"] = k
        ws_inputs[f"B{r}"] = v
        ws_inputs[f"A{r}"].border = border_thin
        ws_inputs[f"B{r}"].border = border_thin

    # Accrued & dirty formulas (ACT/365 simplified accrual; replace if you use different convention)
    ws_inputs["B30"].value = "=100*$B$10*($B$19-$B$21)/365"
    ws_inputs["B31"].value = "=100*$B$10*($B$20-$B$21)/365"
    ws_inputs["B32"].value = "=B27+B30"
    ws_inputs["B33"].value = "=B28+B31"

    # Highlight user inputs
    for addr in ["B27", "B28", "B29"]:
        ws_inputs[addr].fill = input_fill
        ws_inputs[addr].number_format = "0.0000"
    for addr in ["B30", "B31", "B32", "B33"]:
        ws_inputs[addr].number_format = "0.00000"

    # -----------------------------
    # Curve nodes 1/2/3/5/7/10 and KRDV01
    # -----------------------------
    ws_inputs["A35"] = "China Sovereign CNH Yield Curve Nodes (1/2/3/5/7/10) + Key-Rate DV01s"
    ws_inputs["A35"].font = section_font

    curve_headers = ["Tenor (Y)", "Yield t0", "Yield t1", "ΔYield (bp)", "KRDV01 (pts/1bp)", "Notes"]
    for col, h in enumerate(curve_headers, start=1):
        c = ws_inputs.cell(row=36, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border_thin

    tenors = [1, 2, 3, 5, 7, 10]
    # Illustrative yields and KRDV01; replace with pricer outputs
    y0 = [0.0205, 0.0210, 0.0215, 0.0225, 0.0235, 0.0250]
    y1 = [0.0202, 0.0207, 0.0212, 0.0223, 0.0234, 0.0249]
    krdv01 = [0.004, 0.008, 0.012, 0.015, 0.006, 0.001]

    for i, tnr in enumerate(tenors):
        r = 37 + i
        ws_inputs.cell(r, 1, tnr).border = border_thin
        ws_inputs.cell(r, 2, y0[i]).border = border_thin
        ws_inputs.cell(r, 3, y1[i]).border = border_thin
        ws_inputs.cell(r, 4).value = f"=(C{r}-B{r})*10000"  # ΔYield in bp
        ws_inputs.cell(r, 4).border = border_thin
        ws_inputs.cell(r, 5, krdv01[i]).border = border_thin
        ws_inputs.cell(r, 6, "Input node & KRDV01 from your pricer").border = border_thin

        ws_inputs.cell(r, 2).number_format = "0.0000%"
        ws_inputs.cell(r, 3).number_format = "0.0000%"
        ws_inputs.cell(r, 4).number_format = "0.0"
        ws_inputs.cell(r, 5).number_format = "0.0000"

        # user inputs highlighted
        ws_inputs.cell(r, 2).fill = input_fill
        ws_inputs.cell(r, 3).fill = input_fill
        ws_inputs.cell(r, 5).fill = input_fill

    # DV01 and SDV01
    ws_inputs["A44"] = "DV01 checks / curve-only inputs"
    ws_inputs["A44"].font = section_font

    ws_inputs["A45"] = "Sum of KRDV01 (DV01 approx, pts/1bp)"
    ws_inputs["B45"].value = "=SUM(E37:E42)"
    ws_inputs["B45"].number_format = "0.0000"
    ws_inputs["A45"].border = border_thin
    ws_inputs["B45"].border = border_thin

    ws_inputs["A46"] = "Spread DV01 (SDV01, pts/1bp)"
    ws_inputs["B46"] = 0.047
    ws_inputs["B46"].number_format = "0.0000"
    ws_inputs["B46"].fill = input_fill
    ws_inputs["A46"].border = border_thin
    ws_inputs["B46"].border = border_thin

    # -----------------------------
    # Z-spread inputs + bucket lookup (Rating+Sector+Tenor)
    # -----------------------------
    ws_inputs["A48"] = "Z-Spread Inputs (over China sovereign CNH curve)"
    ws_inputs["A48"].font = section_font

    z_rows = [
        ("Bond Z-spread at t0 (bp)", 120),
        ("Bond Z-spread at t1 (bp)", 145),
        ("Bond ΔZ-spread (bp)", None),
        ("Bucket key (Rating|Sector|TenorY)", None),
        ("Class Z-spread at t0 (bp) [lookup]", None),
        ("Class Z-spread at t1 (bp) [lookup]", None),
        ("Class ΔZ-spread (bp)", None),
        ("Idiosyncratic ΔZ (bond - class, bp)", None),
    ]
    start_row = 49
    for i, (k, v) in enumerate(z_rows):
        r = start_row + i
        ws_inputs[f"A{r}"] = k
        ws_inputs[f"B{r}"] = v
        ws_inputs[f"A{r}"].border = border_thin
        ws_inputs[f"B{r}"].border = border_thin

    # Derived Z fields
    ws_inputs["B51"].value = "=B50*1-B49*1"
    ws_inputs["B52"].value = '=$B$14&"|"&$B$15&"|"&$B$16'

    # Lookup class Z from Bucket_Data using VLOOKUP (non-dynamic-array)
    ws_inputs["B53"].value = '=IFERROR(VLOOKUP($B$52,Bucket_Data!$A$3:$D$200,2,FALSE),"")'
    ws_inputs["B54"].value = '=IFERROR(VLOOKUP($B$52,Bucket_Data!$A$3:$D$200,3,FALSE),"")'

    ws_inputs["B55"].value = "=B54*1-B53*1"
    ws_inputs["B56"].value = "=B51*1-B55*1"

    for addr in ["B49", "B50"]:
        ws_inputs[addr].fill = input_fill
        ws_inputs[addr].number_format = "0"
    for addr in ["B51", "B53", "B54", "B55", "B56"]:
        ws_inputs[addr].number_format = "0"

    # Layout
    ws_inputs.column_dimensions["A"].width = 52
    ws_inputs.column_dimensions["B"].width = 28
    ws_inputs.column_dimensions["C"].width = 14
    ws_inputs.column_dimensions["D"].width = 14
    ws_inputs.column_dimensions["E"].width = 18
    ws_inputs.column_dimensions["F"].width = 38
    ws_inputs.freeze_panes = "A4"

    # -----------------------------
    # Attribution
    # -----------------------------
    ws_attr["A1"] = "P&L Attribution (Z-spread over China sovereign CNH curve)"
    ws_attr["A1"].font = title_font
    ws_attr["A3"] = "All amounts are shown per 100 par (price points) and then scaled to CNH P&L using Notional."
    ws_attr["A3"].alignment = wrap_left

    attr_headers = ["Component", "Formula (per 100)", "Value (per 100)", "Notes"]
    for col, h in enumerate(attr_headers, start=1):
        c = ws_attr.cell(row=5, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border_thin

    # Main attribution (per 100)
    main_rows = [
        ("Dirty price t0", "=Inputs!B32", "=Inputs!B32", "From Inputs."),
        ("Dirty price t1", "=Inputs!B33", "=Inputs!B33", "From Inputs."),
        ("Coupon cash (per 100)", "=Inputs!B29", "=Inputs!B29", "Cash coupon paid during period (if any)."),
        ("Total P&L (per 100)", "=(Inputs!B33-Inputs!B32)+Inputs!B29", "=(Inputs!B33-Inputs!B32)+Inputs!B29",
         "ΔDirty + coupon cash."),
        ("Income / Carry", "=(Inputs!B31-Inputs!B30)+Inputs!B29", "=(Inputs!B31-Inputs!B30)+Inputs!B29",
         "ΔAccrued + coupon cash."),
        ("Base curve effect", "=-SUMPRODUCT(Inputs!E37:E42,Inputs!D37:D42)",
         "=-SUMPRODUCT(Inputs!E37:E42,Inputs!D37:D42)", "KRDV01 × node yield changes (bp)."),
        ("Class spread effect", "=-Inputs!B46*$C$24", "=-Inputs!B46*$C$24", "SDV01 × class ΔZ (bp)."),
        ("Selection (idiosyncratic)", "=-Inputs!B46*$C$25", "=-Inputs!B46*$C$25", "SDV01 × (bond ΔZ - class ΔZ)."),
        ("Explained subtotal", "=SUM(C10:C13)", "=SUM(C10:C13)", "Income + base + class + selection."),
        ("Residual", "=C9-C14", "=C9-C14", "Total - explained."),
    ]
    for i, (comp, ftxt, fval, note) in enumerate(main_rows):
        r = 6 + i
        ws_attr.cell(r, 1, comp).border = border_thin
        ws_attr.cell(r, 2, ftxt).border = border_thin
        ws_attr.cell(r, 3).value = fval
        ws_attr.cell(r, 3).border = border_thin
        ws_attr.cell(r, 4, note).border = border_thin
        ws_attr.cell(r, 3).number_format = "0.00000"

    # Diagnostics: Z-spreads (bp)
    ws_attr["A17"] = "Z-spread diagnostics (bp)"
    ws_attr["A17"].font = section_font

    diag_headers = ["Item", "Formula", "Value", "Notes"]
    for col, h in enumerate(diag_headers, start=1):
        c = ws_attr.cell(row=18, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border_thin

    diag_rows = [
        ("Bond Z t0 (bp)", "=Inputs!B49", "=Inputs!B49", "Bond Z-spread at t0 (input)."),
        ("Bond Z t1 (bp)", "=Inputs!B50", "=Inputs!B50", "Bond Z-spread at t1 (input)."),
        ("Bond ΔZ (bp)", "=Inputs!B50-Inputs!B49", "=Inputs!B50-Inputs!B49", "Bond spread move."),
        ("Class Z t0 (bp)", "=Inputs!B53", "=Inputs!B53", "Class Z t0 from bucket lookup."),
        ("Class Z t1 (bp)", "=Inputs!B54", "=Inputs!B54", "Class Z t1 from bucket lookup."),
        ("Class ΔZ (bp)", "=C23-C22", "=C23-C22", "Class spread move."),
        ("Idiosyncratic ΔZ (bp)", "=C21-C24", "=C21-C24", "Bond ΔZ minus class ΔZ."),
        ("Computed bucket key (Rating|Sector|TenorY)", '=Inputs!B14&"|"&Inputs!B15&"|"&Inputs!B16',
         '=Inputs!B14&"|"&Inputs!B15&"|"&Inputs!B16', "Diagnostic only."),
        ("Bucket key exists in Bucket_Data?", '=IF(ISNUMBER(MATCH(C26,Bucket_Data!$A$3:$A$200,0)),"OK","CHECK")',
         '=IF(ISNUMBER(MATCH(C26,Bucket_Data!$A$3:$A$200,0)),"OK","CHECK")', "Diagnostic only."),
    ]
    for i, (item, ftxt, fval, note) in enumerate(diag_rows):
        r = 19 + i
        ws_attr.cell(r, 1, item).border = border_thin
        ws_attr.cell(r, 2, ftxt).border = border_thin
        ws_attr.cell(r, 3).value = fval
        ws_attr.cell(r, 3).border = border_thin
        ws_attr.cell(r, 4, note).border = border_thin
        if "bp" in item.lower():
            ws_attr.cell(r, 3).number_format = "0"

    # Scaling to CNH P&L
    ws_attr["A28"] = "Scaling to CNH P&L"
    ws_attr["A28"].font = section_font

    scale_rows = [
        ("Notional (face, CNH)", "=Inputs!B23", "=Inputs!B23", "Face amount."),
        ("Total P&L (CNH)", "'=(Notional/100)*TotalPnL_per100 - FundingCost", "=(Inputs!B23/100)*C9-Inputs!B24",
         "Scaled total; minus funding."),
        ("Explained P&L (CNH)", "'=(Notional/100)*Explained_per100", "=(Inputs!B23/100)*C14",
         "Scaled explained subtotal."),
        ("Residual P&L (CNH)", "'=(Notional/100)*Residual_per100", "=(Inputs!B23/100)*C15", "Scaled residual."),
    ]
    for i, (item, ftxt, fval, note) in enumerate(scale_rows):
        r = 29 + i
        ws_attr.cell(r, 1, item).border = border_thin
        ws_attr.cell(r, 2, ftxt).border = border_thin
        ws_attr.cell(r, 3).value = fval
        ws_attr.cell(r, 3).border = border_thin
        ws_attr.cell(r, 4, note).border = border_thin
        ws_attr.cell(r, 3).number_format = "#,##0"

    # Reconciliation diagnostics (why Total vs Explained differs)
    ws_attr["A34"] = "Reconciliation diagnostics (why Total vs Explained differs)"
    ws_attr["A34"].font = section_font

    recon_rows = [
        ("Implied spread P&L from Total (per 100)", "'=Total - Income - Base", "=C9-C10-C11", None),
        ("Implied total ΔZ (bp) from price move", "'=-(ImpliedSpreadPnL)/SDV01", "=-(C35)/Inputs!B46", None),
        ("Input Bond ΔZ (bp)", "'=BondZ_t1 - BondZ_t0", "=C21", None),
        ("Bond ΔZ minus implied ΔZ (bp)", "'=BondΔZ - ImpliedΔZ", "=C37-C36", None),
        ("Implied parallel curve shift (bp) needed if Bond/Class ΔZ are correct",
         "'=-(Total - Income - SpreadTotal)/DV01",
         "=-(C9-C10-(C12+C13))/Inputs!B45", None),
        ("Interpretation", None,
         "If Bond ΔZ is far from implied ΔZ, your price marks and Z-spreads are inconsistent (or units/scaling mismatch).",
         None),
    ]
    for i, (item, ftxt, fval, note) in enumerate(recon_rows):
        r = 35 + i
        ws_attr.cell(r, 1, item).border = border_thin
        ws_attr.cell(r, 2, "" if ftxt is None else ftxt).border = border_thin
        ws_attr.cell(r, 3).value = fval
        ws_attr.cell(r, 3).border = border_thin
        ws_attr.cell(r, 4, "" if note is None else note).border = border_thin

        if r == 35:
            ws_attr.cell(r, 3).number_format = "0.00000"
        if r in (36, 37, 38, 39):
            ws_attr.cell(r, 3).number_format = "0.0"
        if r == 40:
            ws_attr.cell(r, 3).number_format = "@"

    # Layout
    ws_attr.column_dimensions["A"].width = 40
    ws_attr.column_dimensions["B"].width = 46
    ws_attr.column_dimensions["C"].width = 30
    ws_attr.column_dimensions["D"].width = 52
    ws_attr.freeze_panes = "A6"

    # Save
    wb.save(out_path)


if __name__ == "__main__":
    build_dimsum_pnl_attribution_xlsx("DimSum_PnL_Attribution_HK0001241626_v2_openpyxl.xlsx")
    print("Saved: DimSum_PnL_Attribution_HK0001241626_v2_openpyxl.xlsx")
