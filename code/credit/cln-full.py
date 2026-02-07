from __future__ import annotations

from dataclasses import dataclass
from datetime import date
import math
from typing import List, Dict, Tuple, Any

from openpyxl import Workbook


# ==============================================================================
# 0) Small utilities
# ==============================================================================

def d(iso: str) -> date:
    """Parse YYYY-MM-DD into datetime.date"""
    return date.fromisoformat(iso)


def act365(a: date, b: date) -> float:
    """ACT/365F year fraction"""
    return (b - a).days / 365.0


def fmt(x: float, nd: int = 2) -> str:
    return f"{x:,.{nd}f}"


def fmt_pct(x: float, nd: int = 4) -> str:
    return f"{x * 100:.{nd}f}%"


def fmt_bp(x: float, nd: int = 2) -> str:
    return f"{x * 10000:.{nd}f} bp"


def print_rows(title: str, headers: List[str], rows: List[List[Any]], max_rows: int = 999) -> None:
    print("\n" + title)
    print("-" * len(title))
    print(" | ".join(headers))
    print("-" * (len(" | ".join(headers))))
    for i, r in enumerate(rows[:max_rows]):
        print(" | ".join(str(v) for v in r))
    if len(rows) > max_rows:
        print(f"... ({len(rows) - max_rows} more rows)")


# ==============================================================================
# 1) Trade + dummy market data
# ==============================================================================

# ---- Trade terms (dummy but consistent) ----
N = 102_000_000  # Notional
c = 0.1062  # Fixed coupon (10.62% p.a.)
R = 0.40  # Recovery assumption for reference credit leg

t0 = d("2023-09-29")  # Previous business day
t1 = d("2023-10-02")  # Your requested valuation date

last_coupon = d("2023-08-14")
maturity = d("2025-11-12")

# ---- Coupon schedule (quarterly) ----
pay_dates = list(map(d, [
    "2023-11-13", "2024-02-12", "2024-05-13", "2024-08-12",
    "2024-11-12", "2025-02-12", "2025-05-12", "2025-08-12",
    "2025-11-12"  # maturity
]))

start_dates = [last_coupon] + pay_dates[:-1]  # period start = last pay date

# ---- Dummy curve nodes (tenor in years) ----
tenors = [0.25, 0.50, 1.00, 2.00, 3.00, 5.00]

swap_t0 = [0.0830, 0.0840, 0.0850, 0.0860, 0.0870, 0.0880]
swap_t1 = [0.0836, 0.0846, 0.0856, 0.0866, 0.0876, 0.0886]  # +6bp shift

iss_t0 = [0.0140, 0.0145, 0.0150, 0.0155, 0.0160, 0.0165]
iss_t1 = [0.0142, 0.0147, 0.0152, 0.0157, 0.0162, 0.0167]  # +2bp shift

ref_t0 = [0.0200, 0.0205, 0.0210, 0.0220, 0.0230, 0.0240]
ref_t1 = [0.0188, 0.0193, 0.0198, 0.0208, 0.0218, 0.0228]  # -12bp tightening


# ==============================================================================
# 2) Linear interpolation (explicit formula)
# ==============================================================================

def interp_linear(x: float, xs: List[float], ys: List[float]) -> Tuple[float, str]:
    """Linear interpolation with flat extrapolation."""
    if x <= xs[0]:
        return ys[0], f"y = ys[0] (flat extrapolate left because x={x:.6f} <= {xs[0]:.2f})"
    if x >= xs[-1]:
        return ys[-1], f"y = ys[-1] (flat extrapolate right because x={x:.6f} >= {xs[-1]:.2f})"

    # find bracketing indices (xs sorted asc)
    lo = 0
    hi = len(xs) - 1
    while hi - lo > 1:
        mid = (lo + hi) // 2
        if xs[mid] <= x:
            lo = mid
        else:
            hi = mid

    x0, x1 = xs[lo], xs[hi]
    y0, y1 = ys[lo], ys[hi]
    y = y0 + (y1 - y0) * (x - x0) / (x1 - x0)
    f = (
        f"y = y0 + (y1-y0)*(x - x0) / (x1 - x0) "
        f"where x={x:.6f}, x0={x0:.2f}, x1={x1:.2f}, y0={y0:.6f}, y1={y1:.6f}"
    )
    return y, f


# ==============================================================================
# 3) Build cashflow schedule (explicit formulas)
# ==============================================================================

schedule_rows: List[Dict[str, Any]] = []

for sd, pd in zip(start_dates, pay_dates):
    days = (pd - sd).days
    yearfrac = days / 365.0

    # Formula strings (Excel-style)
    f_days = f"Days = PayDate - StartDate = ({pd.isoformat()}) - ({sd.isoformat()})"
    f_yf = f"YearFrac = Days/365 = {days}/365"

    # Coupon CF = N * c * YearFrac
    coup_cf = N * c * yearfrac
    f_coup = f"CouponCF = N*c*YearFrac = {N}*{c}*{yearfrac:.10f}"

    # Principal CF = N if PayDate == maturity else 0
    prin_cf = N if pd == maturity else 0.0
    f_prin = f"PrincipalCF = IF(PayDate==Maturity, N, 0)"

    total_cf = coup_cf + prin_cf
    f_total = "TotalCF = CouponCF + PrincipalCF"

    schedule_rows.append({
        "StartDate": sd,
        "PayDate": pd,
        "Days": days, "Days_formula": f_days,
        "YearFrac": yearfrac, "YearFrac_formula": f_yf,
        "CouponCF": coup_cf, "CouponCF_formula": f_coup,
        "PrincipalCF": prin_cf, "PrincipalCF_formula": f_prin,
        "TotalCF": total_cf, "TotalCF_formula": f_total,
    })


# ==============================================================================
# 4) Valuation building blocks (explicit formulas)
# ==============================================================================

def df_issuer(asof: date, pay: date,
              swap_curve: List[float],
              issuer_spread_curve: List[float]) -> Tuple[float, Dict[str, str]]:
    """Continuous-comp issuer DF: DF = exp(-(r+sI)*tau)"""
    tau = act365(asof, pay)
    f_tau = f"Tau = (PayDate - AsOf)/365 = ({pay.isoformat()} - {asof.isoformat()})/365"

    r, f_r = interp_linear(tau, tenors, swap_curve)
    sI, f_sI = interp_linear(tau, tenors, issuer_spread_curve)

    df = math.exp(-(r + sI) * tau)
    f_df = f"DF = exp(-(r+sI)*Tau) = exp(-({r:.10f}+{sI:.10f})*{tau:.10f})"

    return df, {
        "Tau": f_tau,
        "SwapRate": f_r,
        "IssuerSpr": f_sI,
        "DF": f_df
    }


def pv_note(asof: date,
            swap_curve: List[float],
            issuer_spread_curve: List[float]) -> Tuple[float, List[Dict[str, Any]]]:
    """PV_note(asof) = sum_{pay>asof} TotalCF_i * DF_issuer(asof, pay_i)"""
    pv = 0.0
    detail = []

    for row in schedule_rows:
        pd = row["PayDate"]
        if pd <= asof:
            continue

        df, f = df_issuer(asof, pd, swap_curve, issuer_spread_curve)
        pv_contrib = row["TotalCF"] * df
        f_pv = f"PV_contrib = TotalCF * DF = {row['TotalCF']:.6f} * {df:.10f}"

        pv += pv_contrib

        detail.append({
            "PayDate": pd.isoformat(),
            "TotalCF": row["TotalCF"],
            "Tau": act365(asof, pd),
            "DF": df,
            "PV_contrib": pv_contrib,
            "formula_Tau": f["Tau"],
            "formula_SwapRate": f["SwapRate"],
            "formula_IssuerSpr": f["IssuerSpr"],
            "formula_DF": f["DF"],
            "formula_PV": f_pv
        })

    return pv, detail


# ==============================================================================
# 5) Credit leg (CDS-style proxy) with explicit formulas
# ==============================================================================

def cds_proxy(asof: date,
              swap_curve: List[float],
              issuer_spread_curve: List[float],
              ref_spread_curve: List[float],
              S_fixed: float) -> Tuple[float, float, float, List[Dict[str, Any]]]:
    """
    CDS-like proxy (buyer of protection):
    For each future period i:
      delta_i = (PayDate - max(StartDate, AsOf))/365
      tau_i   = (PayDate - AsOf)/365
      z_i     = interp(ref_spread_curve at tau_i)
      h_i     = z_i / (1-R)
      Q_end   = Q_prev * exp(-h_i * delta_i)
      DF_i    = issuer discount factor (same as note discounting here)
      ProtContrib_i = (1-R)*N*DF_i*(Q_prev - Q_end)
      RPV01Contrib_i = N*DF_i*delta_i*Q_end
    ProtPV = sum ProtContrib
    RPV01  = sum RPV01Contrib
    PV_CDS_buy = ProtPV - S_fixed * RPV01

    Investor in a typical CLN is short protection:
      PV_credit_investor = - PV_CDS_buy
    """
    Q_prev = 1.0
    protPV = 0.0
    rpv01 = 0.0
    detail = []

    for sd, pd in zip(start_dates, pay_dates):
        if pd <= asof:
            continue

        start_eff = max(sd, asof)
        delta = act365(start_eff, pd)  # accrual length from asof within this period
        tau = act365(asof, pd)  # time to payment

        # z(tau) from curve
        z, f_z = interp_linear(tau, tenors, ref_spread_curve)

        # hazard
        h = z / (1.0 - R)
        f_h = f"h = z/(1-R) = {z:.10f}/(1-{R:.2f})"

        # survival update
        Q_end = Q_prev * math.exp(-h * delta)
        f_Q = f"Q_end = Q_prev * exp(-h*delta) = {Q_prev:.10f}*exp(-{h:.10f}*{delta:.10f})"

        # discount factor (issuer curve)
        DF, f_df = df_issuer(asof, pd, swap_curve, issuer_spread_curve)

        # protection contribution
        prot_contrib = (1.0 - R) * N * DF * (Q_prev - Q_end)
        f_prot = (
            f"ProtContrib = (1-R)*N*DF*(Q_prev-Q_end) ",
            f"= {1 - R:.2f}*{N}*{DF:.10f}*({Q_prev:.10f}-{Q_end:.10f})"
        )

        # premium PV01 contribution (RPV01)
        rpv01_contrib = N * DF * delta * Q_end
        f_rpv01 = (
            f"RPV01Contrib = N*DF*delta*Q_end ",
            f"= {N}*{DF:.10f}*{delta:.10f}*{Q_end:.10f}"
        )

        protPV += prot_contrib
        rpv01 += rpv01_contrib

        detail.append({
            "StartDate": sd.isoformat(),
            "PayDate": pd.isoformat(),
            "delta": delta,
            "tau": tau,
            "z": z,
            "h": h,
            "Q_prev": Q_prev,
            "Q_end": Q_end,
            "DF": DF,
            "ProtContrib": prot_contrib,
            "RPV01Contrib": rpv01_contrib,
            "formula_delta": f"delta = (PayDate - max(StartDate,AsOf))/365 = ({pd.isoformat()} - max({sd.isoformat()},{asof.isoformat()}))/365",
            "formula_tau": f"tau = (PayDate - AsOf)/365 = ({pd.isoformat()} - {asof.isoformat()})/365",
            "formula_z": f_z,
            "formula_h": f_h,
            "formula_Q": f_Q,
            "formula_DF": f_df["DF"],
            "formula_ProtContrib": f_prot,
            "formula_RPV01Contrib": f_rpv01
        })

        Q_prev = Q_end

    pv_cds_buy = protPV - S_fixed * rpv01
    f_pv = f"PV_CDS_buy = ProtPV - S_fixed*RPV01 = {protPV:.6f} - {S_fixed:.10f}*{rpv01:.6f}"

    # We return the PV and also the components so you can see calibration etc.
    return pv_cds_buy, protPV, rpv01, detail


def calibrate_S_fixed_par_at_t0() -> Tuple[float, float, float]:
    """
    Par spread calibration at t0:
      S_fixed = ProtPV(t0) / RPV01(t0)
    computed using CDS proxy with S_fixed=0 to get ProtPV and RPV01.
    """
    pv_cds_buy_0, protPV0, rpv010, _ = cds_proxy(
        asof=t0,
        swap_curve=swap_t0,
        issuer_spread_curve=iss_t0,
        ref_spread_curve=ref_t0,
        S_fixed=0.0
    )
    S = protPV0 / rpv010
    return S, protPV0, rpv010


# ==============================================================================
# 6) Scenario valuations and stepwise attribution
# ==============================================================================

S_fixed, protPV0, rpv010 = calibrate_S_fixed_par_at_t0()


def pv_total_investor(asof: date,
                      swap_curve: List[float],
                      issuer_spread_curve: List[float],
                      ref_spread_curve: List[float]) -> Dict[str, Any]:
    """
    Investor view for CLN decomposition:
      PV_total = PV_note + PV_credit
      PV_credit(investor) = - PV_CDS_buy_protection (since investor sells protection)

    Returns dict with totals and per-leg detail.
    """
    notePV, note_detail = pv_note(asof, swap_curve, issuer_spread_curve)

    pv_cds_buy, protPV, rpv01, cds_detail = cds_proxy(
        asof=asof,
        swap_curve=swap_curve,
        issuer_spread_curve=issuer_spread_curve,
        ref_spread_curve=ref_spread_curve,
        S_fixed=S_fixed
    )

    creditPV = -pv_cds_buy
    totalPV = notePV + creditPV

    return {
        "asof": asof.isoformat(),
        "PV_note": notePV,
        "PV_cds_buy": pv_cds_buy,
        "PV_credit_investor": creditPV,
        "PV_total_investor": totalPV,
        "ProtPV": protPV,
        "RPV01": rpv01,
        "note_detail": note_detail,
        "cds_detail": cds_detail,
    }


# --- PV ladder scenarios ---
# 1) PV_t0: asof t0, all curves t0
sc_t0 = pv_total_investor(t0, swap_t0, iss_t0, ref_t0)

# 2) PV_theta: asof t1, curves frozen at t0
sc_theta = pv_total_investor(t1, swap_t0, iss_t0, ref_t0)

# 3) PV_rates: asof t1, swap moved to t1, issuer/ref still t0
sc_rates = pv_total_investor(t1, swap_t1, iss_t0, ref_t0)

# 4) PV_issuer: asof t1, swap t1, issuer moved to t1, ref still t0
sc_issuer = pv_total_investor(t1, swap_t1, iss_t1, ref_t0)

# 5) PV_ref: asof t1, swap t1, issuer t1, ref moved to t1
sc_ref = pv_total_investor(t1, swap_t1, iss_t1, ref_t1)

PV_t0 = sc_t0["PV_total_investor"]
PV_theta = sc_theta["PV_total_investor"]
PV_rates = sc_rates["PV_total_investor"]
PV_issuer = sc_issuer["PV_total_investor"]
PV_ref = sc_ref["PV_total_investor"]

# No coupon payments between 2023-09-29 and 2023-10-02 in this schedule => cashflows=0.
carry = PV_theta - PV_t0
rates = PV_rates - PV_theta
issuer = PV_issuer - PV_rates
ref = PV_ref - PV_issuer
total = PV_ref - PV_t0
resid = total - (carry + rates + issuer + ref)

# ==============================================================================
# 7) Create an in-memory workbook (openpyxl) WITHOUT saving
# ==============================================================================

wb = Workbook()

# Inputs sheet
ws = wb.active
ws.title = "Inputs"
ws.append(["Field", "Value", "Formula/Comment"])
ws.append(["Notional N", N, "Input"])
ws.append(["Coupon c", c, "Input"])
ws.append(["Recovery R", R, "Input"])
ws.append(["t0", t0.isoformat(), "Input"])
ws.append(["t1", t1.isoformat(), "Input"])
ws.append(["LastCoupon", last_coupon.isoformat(), "Input"])
ws.append(["Maturity", maturity.isoformat(), "Input"])
ws.append(["S_fixed (par @ t0)", S_fixed, "S_fixed = ProtPV0 / RPV010"])
ws.append(["ProtPV(t0) (S=0)", protPV0, "ProtPV from CDS proxy at t0"])
ws.append(["RPV01(t0)", rpv010, "RPV01 from CDS proxy at t0"])

# Curves sheet
ws = wb.create_sheet("Curves")
ws.append(["Tenor", "Swap_t0", "Swap_t1", "IssuerSpr_t0", "IssuerSpr_t1", "RefSpr_t0", "RefSpr_t1"])
for i in range(len(tenors)):
    ws.append([tenors[i], swap_t0[i], swap_t1[i], iss_t0[i], iss_t1[i], ref_t0[i], ref_t1[i]])

# Schedule sheet with values + formula strings
ws = wb.create_sheet("Schedule")
ws.append([
    "StartDate", "PayDate",
    "Days", "Days_formula",
    "YearFrac", "YearFrac_formula",
    "CouponCF", "CouponCF_formula",
    "PrincipalCF", "PrincipalCF_formula",
    "TotalCF", "TotalCF_formula"
])
for r in schedule_rows:
    ws.append([
        r["StartDate"].isoformat(), r["PayDate"].isoformat(),
        r["Days"], r["Days_formula"],
        r["YearFrac"], r["YearFrac_formula"],
        r["CouponCF"], r["CouponCF_formula"],
        r["PrincipalCF"], r["PrincipalCF_formula"],
        r["TotalCF"], r["TotalCF_formula"]
    ])

# Note valuation detail sheet (for PV_ref scenario as an example)
ws = wb.create_sheet("NoteDetail_PV_ref")
ws.append([
    "PayDate", "TotalCF", "Tau", "DF", "PV_contrib",
    "formula_Tau", "formula_SwapRate", "formula_IssuerSpr", "formula_DF", "formula_PV"
])
for r in sc_ref["note_detail"]:
    ws.append([
        r["PayDate"], r["TotalCF"], r["Tau"], r["DF"], r["PV_contrib"],
        r["formula_Tau"], r["formula_SwapRate"], r["formula_IssuerSpr"], r["formula_DF"], r["formula_PV"]
    ])

# Credit valuation detail (for PV_ref scenario)
ws = wb.create_sheet("CreditDetail_PV_ref")
ws.append([
    "StartDate", "PayDate", "delta", "tau", "z", "h", "Q_prev", "Q_end", "DF",
    "ProtContrib", "RPV01Contrib",
    "f_delta", "f_tau", "f_z", "f_h", "f_Q", "f_DF", "f_ProtContrib", "f_RPV01Contrib"
])
for r in sc_ref["cds_detail"]:
    ws.append([
        r["StartDate"], r["PayDate"], r["delta"], r["tau"], r["z"], r["h"],
        r["Q_prev"], r["Q_end"], r["DF"],
        r["ProtContrib"], r["RPV01Contrib"],
        r["formula_delta"], r["formula_tau"], r["formula_z"], r["formula_h"],
        r["formula_Q"], r["formula_DF"], r["formula_ProtContrib"][1], r["formula_RPV01Contrib"][1]
    ])

# PV ladder sheet
ws = wb.create_sheet("PV Ladder")
ws.append(["Scenario", "PV_total", "PV_note", "PV_credit", "Comment"])
ws.append(["PV_t0", PV_t0, sc_t0["PV_note"], sc_t0["PV_credit_investor"], "asof=t0, swap/issuer/ref=t0"])
ws.append(["PV_theta", PV_theta, sc_theta["PV_note"], sc_theta["PV_credit_investor"], "asof=t1, curves frozen @t0"])
ws.append(["PV_rates", PV_rates, sc_rates["PV_note"], sc_rates["PV_credit_investor"], "swap->t1 only"])
ws.append(
    ["PV_issuer", PV_issuer, sc_issuer["PV_note"], sc_issuer["PV_credit_investor"], "issuer->t1 (swap already t1)"])
ws.append(["PV_ref", PV_ref, sc_ref["PV_note"], sc_ref["PV_credit_investor"], "ref->t1 (swap, issuer already t1)"])

# Attribution sheet with formulas
ws = wb.create_sheet("Attribution")
ws.append(["Component", "Formula", "Value"])
ws.append(["Carry/Theta", "PV_theta - PV_t0", carry])
ws.append(["Rates", "PV_rates - PV_theta", rates])
ws.append(["Issuer spread", "PV_issuer - PV_rates", issuer])
ws.append(["Reference spread", "PV_ref - PV_issuer", ref])
ws.append(["Total", "PV_ref - PV_t0", total])
ws.append(["Residual", "Total - (Carry+Rates+Issuer+Ref)", resid])

#wb.save('ZAR_CLN_PnL_Attribution_20231002_full.xlsx')

# ==============================================================================
# 8) Print outputs (values + formulas)
# ==============================================================================

print("\n" + "=" * 60)
print("CLN P&L Attribution (No Default)")
print(f"AsOf t0 = {t0.isoformat()} -> t1 = {t1.isoformat()}")
print("=" * 60)

print("\n--- Calibration at t0 ---")
print(f"S_fixed (par @ t0) = {fmt_pct(S_fixed, 6)} = {fmt_bp(S_fixed)}")
print(f"Formula: S_fixed = ProtPV(t0) / RPV01(t0)")
print(f"ProtPV(t0) = {fmt(protPV0)}")
print(f"RPV01(t0)  = {fmt(rpv010)}")

# Print schedule (compact)
sched_print = []
for r in schedule_rows:
    sched_print.append([
        r["StartDate"].isoformat(),
        r["PayDate"].isoformat(),
        r["Days"], f"{r['YearFrac']:.10f}",
        fmt(r["TotalCF"]),
        r["YearFrac_formula"],
        r["CouponCF_formula"],
        r["PrincipalCF_formula"]
    ])

print_rows(
    title="Cashflow Schedule (with key formulas)",
    headers=["StartDate", "PayDate", "Days", "YearFrac", "TotalCF", "f_YearFrac", "f_CouponCF", "f_PrincipalCF"],
    rows=sched_print,
    max_rows=20
)

# PV ladder print
ladder_rows = [
    ["PV_t0", fmt(PV_t0), "PV(t0; swap0, issuer0, ref0)"],
    ["PV_theta", fmt(PV_theta), "PV(t1; swap0, issuer0, ref0)"],
    ["PV_rates", fmt(PV_rates), "PV(t1; swap1, issuer0, ref0)"],
    ["PV_issuer", fmt(PV_issuer), "PV(t1; swap1, issuer1, ref0)"],
    ["PV_ref", fmt(PV_ref), "PV(t1; swap1, issuer1, ref1)"]
]

print_rows(
    title="PV Ladder (Investor view, dirty PV)",
    headers=["Scenario", "PV", "Definition"],
    rows=ladder_rows,
    max_rows=10
)

# Attribution print (with formulas)
attrib_rows = [
    ["Carry/Theta", "PV_theta - PV_t0", fmt(carry)],
    ["Rates", "PV_rates - PV_theta", fmt(rates)],
    ["Issuer spread", "PV_issuer - PV_rates", fmt(issuer)],
    ["Reference spread", "PV_ref - PV_issuer", fmt(ref)],
    ["Total", "PV_ref - PV_t0", fmt(total)],
    ["Residual", "Total - sum(components)", fmt(resid)]
]

print_rows(
    title="Daily P&L Attribution (t0 -> t1)",
    headers=["Component", "Formula", "P&L"],
    rows=attrib_rows,
    max_rows=20
)

# Show a few lines of NOTE detail with formulas (PV_ref scenario)
note_detail_rows = []
for r in sc_ref["note_detail"][:3]:
    note_detail_rows.append([
        r["PayDate"],
        fmt(r["TotalCF"]),
        f"{r['Tau']:.6f}",
        f"{r['DF']:.10f}",
        fmt(r["PV_contrib"]),
        r["formula_DF"],
        r["formula_PV"]
    ])

print_rows(
    title="Note Leg Detail (first 3 cashflows, PV_ref scenario) with formulas",
    headers=["PayDate", "TotalCF", "Tau", "DF", "PV_contrib", "f_DF", "f_PV"],
    rows=note_detail_rows,
    max_rows=10
)

# Show a few lines of CREDIT detail with formulas (PV_ref scenario)
credit_detail_rows = []
for r in sc_ref["cds_detail"][:3]:
    credit_detail_rows.append([
        r["PayDate"],
        f"{r['delta']:.6f}",
        f"{r['tau']:.6f}",
        fmt_bp(r["z"]),
        f"{r['Q_prev']:.8f}",
        f"{r['Q_end']:.8f}",
        fmt(r["ProtContrib"]),
        fmt(r["RPV01Contrib"]),
        r["formula_h"],
        r["formula_Q"]
    ])

print_rows(
    title="Credit Leg Detail (first 3 periods, PV_ref scenario) with formulas",
    headers=["PayDate", "delta", "tau", "z", "Q_prev", "Q_end", "ProtContrib", "RPV01Contrib", "f_h", "f_Q"],
    rows=credit_detail_rows,
    max_rows=10
)

print("\nDone. Workbook created in memory via openpyxl, but NOT saved (per your request).")