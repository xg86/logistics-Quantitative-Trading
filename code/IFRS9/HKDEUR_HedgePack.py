"""
HKD/EUR Hedge Pack Generator (from scratch)
-------------------------------------------
Creates an Excel workbook (openpyxl) that contains:
- Trade terms (HKD/EUR CCS)
- Hedged item (EUR shareholder loan; designated EUR layer)
- FVH effectiveness (spot FX component) incl. dollar-offset + regression + sensitivity
- Optional CFH (forecast EUR coupons) incl. FX forwards ladder + effectiveness + regression
- HK holiday list (template) + parameters (dummy spot path / dummy HIBOR)
- CVA parameters + CVA profiles + CVA scorecard
- Journal templates for FVH and CFH (CVA posted through P&L)
All computations are preserved as Excel formulas; Excel will recalculate on open.
Usage:
  python build_hkdeur_hedgepack.py --out HKDEUR_HedgePack.xlsx
"""

from __future__ import annotations

import calendar
import argparse
from datetime import date, timedelta
from typing import List, Set

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName


def build_hk_holidays_2024_2027() -> List[date]:
    """
    Template HK holiday list (replace with official gazetted HK holidays for production).
    """
    hols: List[date] = []
    # 2024 (template)
    hols += [
        date(2024, 1, 1),
        date(2024, 2, 10), date(2024, 2, 12), date(2024, 2, 13),
        date(2024, 3, 29), date(2024, 4, 1), date(2024, 4, 4),
        date(2024, 5, 1), date(2024, 5, 15), date(2024, 6, 10),
        date(2024, 7, 1), date(2024, 9, 18),
        date(2024, 10, 1), date(2024, 10, 11),
        date(2024, 12, 25), date(2024, 12, 26),
    ]
    # 2025 (template)
    hols += [
        date(2025, 1, 1),
        date(2025, 1, 29), date(2025, 1, 30), date(2025, 1, 31),
        date(2025, 4, 4), date(2025, 4, 18), date(2025, 4, 21),
        date(2025, 5, 1), date(2025, 5, 5),
        date(2025, 6, 2), date(2025, 7, 1),
        date(2025, 10, 1), date(2025, 10, 7), date(2025, 10, 29),
        date(2025, 12, 25), date(2025, 12, 26),
    ]
    # 2026 (template)
    hols += [
        date(2026, 1, 1),
        date(2026, 2, 17), date(2026, 2, 18), date(2026, 2, 19),
        date(2026, 4, 3), date(2026, 4, 6), date(2026, 4, 7),
        date(2026, 5, 1),
        date(2026, 5, 25), date(2026, 6, 19),
        date(2026, 7, 1),
        date(2026, 9, 25),
        date(2026, 10, 1), date(2026, 10, 19),
        date(2026, 12, 25), date(2026, 12, 26),
    ]
    # 2027 (provisional template)
    hols += [
        date(2027, 1, 1),
        date(2027, 2, 6), date(2027, 2, 8), date(2027, 2, 9),
        date(2027, 3, 26), date(2027, 3, 29),
        date(2027, 4, 2),
        date(2027, 5, 3), date(2027, 5, 13),
        date(2027, 6, 14),
        date(2027, 7, 1),
        date(2027, 9, 15),
        date(2027, 10, 1), date(2027, 10, 11),
        date(2027, 12, 25), date(2027, 12, 27),
    ]
    return sorted(set(hols))


def is_bd(d: date, holidays: Set[date]) -> bool:
    return d.weekday() < 5 and d not in holidays


def prev_bd(d: date, holidays: Set[date]) -> date:
    while not is_bd(d, holidays):
        d -= timedelta(days=1)
    return d


def modified_following(d: date, holidays: Set[date]) -> date:
    """Move forward to next BD; if that crosses month, move backward."""
    if is_bd(d, holidays):
        return d
    orig_month = d.month
    dd = d
    while not is_bd(dd, holidays):
        dd += timedelta(days=1)
    if dd.month != orig_month:
        dd = d
        while not is_bd(dd, holidays):
            dd -= timedelta(days=1)
    return dd


def add_months(d: date, months: int) -> date:
    m = d.month - 1 + months
    y = d.year + m // 12
    m = m % 12 + 1
    day = min(d.day, calendar.monthrange(y, m)[1])
    return date(y, m, day)


def generate_weekly_obs(start: date, end: date, weekday: int, holidays: Set[date]) -> List[date]:
    """Weekly observation dates on a given weekday, adjusted to previous HK BD."""
    d = start
    while d.weekday() != weekday:
        d += timedelta(days=1)
    obs: List[date] = []
    while d <= end:
        obs.append(prev_bd(d, holidays))
        d += timedelta(days=7)
    return sorted(set(obs))


def generate_roll_dates(start: date, end: date, step_months: int) -> List[date]:
    """Generate roll dates from start stepping by step_months until end; includes end."""
    dates = [start]
    d = start
    while True:
        d = add_months(d, step_months)
        if d > end:
            break
        dates.append(d)
    if dates[-1] != end:
        dates.append(end)
    return dates


def build_hkdeur_hedgepack(output_path: str) -> str:
    """
    Generates the HKD/EUR hedge documentation workbook from scratch (no input file).
    All key calculations are preserved as Excel formulas; Excel will recalc on open.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Styles
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    title_font = Font(bold=True, size=14)
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def add_sheet(name: str):
        return wb.create_sheet(name)

    def set_col_widths(ws, mapping):
        for col, width in mapping.items():
            ws.column_dimensions[col].width = width

    def style_cells(ws, cell_range: str, header: bool = False):
        for row in ws[cell_range]:
            for cell in row:
                cell.border = border
                cell.alignment = center if header else wrap
                if header:
                    cell.font = bold
                    cell.fill = header_fill

    # Holidays
    hols = build_hk_holidays_2024_2027()
    hol_set = set(hols)

    ws_hol = add_sheet("97_HK_Holidays")
    ws_hol["A1"] = "Hong Kong Holidays (Template – replace with official list)"
    ws_hol["A1"].font = title_font
    ws_hol["A3"], ws_hol["B3"] = "Date", "Note"
    style_cells(ws_hol, "A3:B3", header=True)
    for i, d in enumerate(hols, start=4):
        ws_hol[f"A{i}"] = d
        ws_hol[f"A{i}"].number_format = "yyyy-mm-dd"
        ws_hol[f"B{i}"] = "Template holiday"
        style_cells(ws_hol, f"A{i}:B{i}")
    set_col_widths(ws_hol, {"A": 14, "B": 28})
    ws_hol.freeze_panes = "A4"
    wb.defined_names.add(DefinedName("HK_HOL", attr_text=f"'97_HK_Holidays'!$A$4:$A${3+len(hols)}"))

    # Parameters (dummy values)
    ws_p = add_sheet("98_Params")
    ws_p["A1"] = "Parameters"
    ws_p["A1"].font = title_font
    ws_p["A3"], ws_p["B3"], ws_p["C3"] = "Parameter", "Value", "Notes"
    style_cells(ws_p, "A3:C3", header=True)
    params = [
        ("Min_R2", 0.80, "Regression threshold"),
        ("Slope_Lower", 0.80, "Regression slope lower"),
        ("Slope_Upper", 1.25, "Regression slope upper"),
        ("DO_Lower", 0.80, "Dollar offset lower"),
        ("DO_Upper", 1.25, "Dollar offset upper"),
        ("Testing_Frequency", "Weekly", ""),
        ("Obs_Weekday", 4, "Friday = 4"),
        ("Dummy_Spot_Start", 8.4684, "HKD per EUR"),
        ("Dummy_Spot_Weekly_Drift", 0.0010, "0.10% per week"),
        ("Dummy_HIBOR", 0.0400, "4.00% dummy 3M HIBOR"),
        ("PD_1Y", 0.0010, "Dummy"),
        ("Recovery", 0.40, "Dummy"),
        ("DiscountRate", 0.03, "Dummy flat"),
    ]
    for i, (k, v, n) in enumerate(params, start=4):
        ws_p[f"A{i}"], ws_p[f"B{i}"], ws_p[f"C{i}"] = k, v, n
        style_cells(ws_p, f"A{i}:C{i}")
    ws_p["B14"].number_format = "0.00%"
    ws_p["B15"].number_format = "0.00%"
    ws_p["B16"].number_format = "0.00%"
    ws_p["B13"].number_format = "0.00%"
    ws_p["B12"].number_format = "0.0000"
    set_col_widths(ws_p, {"A": 22, "B": 18, "C": 40})
    ws_p.freeze_panes = "A4"

    # Conventions
    ws0 = add_sheet("00_Conventions")
    ws0["A1"] = "Conventions (All cashflows from CPL perspective)"
    ws0["A1"].font = title_font
    ws0["A3"], ws0["B3"] = "Item", "Description"
    style_cells(ws0, "A3:B3", header=True)
    rows = [
        ("Hedged item", "EUR shareholder loan receivable (underlying EUR 20m @ 8.4%, semiannual; designate a EUR layer to match hedge)."),
        ("Hedge instrument", "HKD/EUR CCS (CPL pays EUR fixed, receives HKD float; final exchange: receive HKD notional, pay EUR notional)."),
        ("FV hedge risk", "Spot FX risk (EUR/HKD) on recognised receivable (principal layer + accrued interest)."),
        ("Optional CF hedge", "Spot FX risk on forecast EUR interest receipts (hedged with FX forwards ladder template)."),
        ("Obs dates", "Weekly (Friday) observation dates, adjusted to previous HK business day using HK_HOL."),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        ws0[f"A{i}"], ws0[f"B{i}"] = a, b
        style_cells(ws0, f"A{i}:B{i}")
    set_col_widths(ws0, {"A": 20, "B": 95})
    ws0.freeze_panes = "A4"

    # Trade terms (key cells: C6/C7/C8/C9/C10)
    ws1 = add_sheet("01_TradeTerms")
    ws1["A1"] = "Trade Terms – HKD/EUR CCS"
    ws1["A1"].font = title_font
    ws1["B3"], ws1["C3"] = "Field", "Value"
    style_cells(ws1, "B3:C3", header=True)
    mapping = [
        (4, "Trade reference", "CPL-CCS-202401"),
        (5, "Counterparty", "XYZ Bank"),
        (6, "Effective date", date(2024, 5, 27)),
        (7, "Termination date", date(2027, 5, 27)),
        (8, "EUR notional (CPL pays)", 2_000_000),
        (9, "HKD notional (CPL receives)", 16_936_800),
        (10, "Spot (HKD/EUR) locked for final exchange", "=C9/C8"),
        (11, "EUR fixed rate (%)", 4.25),
        (12, "HKD index", "3M HIBOR"),
        (13, "HKD spread (bps)", 108),
        (14, "BDC", "Modified Following"),
        (15, "Calendar", "Hong Kong (HK_HOL)"),
        (16, "Initial exchange", "No"),
        (17, "Final exchange", "Receive HKD notional / Pay EUR notional"),
    ]
    for r, label, val in mapping:
        ws1[f"B{r}"], ws1[f"C{r}"] = label, val
        style_cells(ws1, f"B{r}:C{r}")
        if isinstance(val, date):
            ws1[f"C{r}"].number_format = "yyyy-mm-dd"
    ws1["C8"].number_format = "#,##0"
    ws1["C9"].number_format = "#,##0"
    ws1["C10"].number_format = "0.0000"
    ws1["C11"].number_format = "0.00"
    set_col_widths(ws1, {"A": 2, "B": 40, "C": 38})
    ws1.freeze_panes = "A4"

    eff = ws1["C6"].value
    term = ws1["C7"].value
    obs = generate_weekly_obs(eff, term, weekday=4, holidays=hol_set)

    # Hedged item
    ws2 = add_sheet("02_HedgedItem")
    ws2["A1"] = "Hedged Item – EUR Shareholder Loan"
    ws2["A1"].font = title_font
    ws2["A3"], ws2["B3"], ws2["C3"] = "Field", "Value", "Notes"
    style_cells(ws2, "A3:C3", header=True)
    fields = [
        (4, "Underlying loan principal (EUR)", 20_000_000, "Underlying exposure (loan agreement)"),
        (5, "Underlying loan fixed rate (%)", 8.4, "p.a."),
        (6, "Coupon frequency", "Semiannual", "20-Feb / 20-Aug per template"),
        (7, "Underlying loan maturity", date(2040, 8, 20), ""),
        (8, "Designated principal layer (EUR)", "=01_TradeTerms!C8", "Designated hedged item layer"),
        (9, "Hedged risk (FVH)", "EUR/HKD spot FX", ""),
        (10, "Hedging instrument reference", "=01_TradeTerms!C4", ""),
        (11, "Hedging instrument EUR notional", "=01_TradeTerms!C8", ""),
        (12, "Hedge ratio (HI/Der)", "=B8/B11", "Expect ~1.00"),
        (13, "FV hedge horizon end date", "=01_TradeTerms!C7", "CCS termination"),
    ]
    for r, f, v, n in fields:
        ws2[f"A{r}"], ws2[f"B{r}"], ws2[f"C{r}"] = f, v, n
        style_cells(ws2, f"A{r}:C{r}")
        if isinstance(v, date):
            ws2[f"B{r}"].number_format = "yyyy-mm-dd"
    ws2["B4"].number_format = "#,##0"
    ws2["B5"].number_format = "0.00"
    ws2["B8"].number_format = "#,##0"
    ws2["B11"].number_format = "#,##0"
    ws2["B12"].number_format = "0.0000"
    set_col_widths(ws2, {"A": 34, "B": 22, "C": 55})
    ws2.freeze_panes = "A4"

    # Hedge designation
    ws3 = add_sheet("03_HedgeDesignation")
    ws3["A1"] = "Hedge Designation (IFRS 9 template)"
    ws3["A1"].font = title_font
    ws3["A3"], ws3["B3"], ws3["C3"] = "Field", "Value", "Notes"
    style_cells(ws3, "A3:C3", header=True)
    hd = [
        (4, "Hedge relationship 1", "Fair Value Hedge (spot FX on EUR loan layer)", ""),
        (5, "Hedged item", "=02_HedgedItem!B8", "Designated EUR principal layer"),
        (6, "Hedging instrument", "=01_TradeTerms!C4", ""),
        (7, "Hedged risk", "EUR/HKD exchange rate risk (fair value changes on principal layer)", ""),
        (8, "Prospective assessment", "Critical terms + Sensitivity + (optional) Regression", ""),
        (9, "Retrospective assessment", "Dollar offset + (optional) Regression", ""),
        (10, "Testing frequency", "=98_Params!B9", "Weekly"),
        (11, "Excluded components", "Non-spot components (rates/basis) – policy dependent", ""),
        (13, "Hedge relationship 2 (optional)", "Cash Flow Hedge (spot FX on forecast EUR coupons)", ""),
        (14, "Hedged item (CFH)", "Forecast EUR interest receipts", ""),
        (15, "Hedging instrument (CFH)", "FX forward ladder", ""),
        (16, "Testing frequency (CFH)", "=98_Params!B9", ""),
    ]
    for r, f, v, n in hd:
        ws3[f"A{r}"], ws3[f"B{r}"], ws3[f"C{r}"] = f, v, n
        style_cells(ws3, f"A{r}:C{r}")
    set_col_widths(ws3, {"A": 30, "B": 55, "C": 45})
    ws3.freeze_panes = "A4"

    # HKD leg schedule
    ws4 = add_sheet("04_Schedule_HKD_Leg")
    ws4["A1"] = "HKD Leg Schedule (CPL receives HKD float)"
    ws4["A1"].font = title_font
    headers = ["Period","AccrualStart","AccrualEnd","PayDate","FixingDate(T-2)","NotionalHKD","Index","Spread(bps)","Fixing(%)","DayCount","DCF","InterestHKD"]
    for c, h in enumerate(headers, 1):
        ws4.cell(3, c, h)
    style_cells(ws4, "A3:L3", header=True)
    set_col_widths(ws4, {get_column_letter(i): w for i, w in enumerate([8,12,12,12,14,16,12,12,10,10,10,16], 1)})
    ws4.freeze_panes = "A4"
    roll_dates = generate_roll_dates(eff, term, 3)
    for i in range(1, len(roll_dates)):
        r = 3 + i
        start = roll_dates[i-1]
        end = roll_dates[i]
        pay = modified_following(end, hol_set)
        ws4[f"A{r}"], ws4[f"B{r}"], ws4[f"C{r}"], ws4[f"D{r}"] = i, start, end, pay
        ws4[f"E{r}"] = f"=WORKDAY.INTL(D{r},-2,\"0000011\",HK_HOL)"
        ws4[f"F{r}"] = "=01_TradeTerms!C9"
        ws4[f"G{r}"] = "=01_TradeTerms!C12"
        ws4[f"H{r}"] = "=01_TradeTerms!C13"
        ws4[f"I{r}"] = "=98_Params!B13"  # dummy 3M HIBOR
        ws4[f"J{r}"] = "ACT/365"
        ws4[f"K{r}"] = f"=(C{r}-B{r})/365"
        ws4[f"L{r}"] = f"=F{r}*((I{r}+H{r}/10000))*K{r}"
        for col in "BCDE":
            ws4[f"{col}{r}"].number_format = "yyyy-mm-dd"
        ws4[f"F{r}"].number_format = "#,##0"
        ws4[f"I{r}"].number_format = "0.00%"
        ws4[f"K{r}"].number_format = "0.0000"
        ws4[f"L{r}"].number_format = "#,##0"
        style_cells(ws4, f"A{r}:L{r}")

    # EUR leg schedule
    ws5 = add_sheet("05_Schedule_EUR_Leg")
    ws5["A1"] = "EUR Leg Schedule (CPL pays fixed EUR)"
    ws5["A1"].font = title_font
    headers = ["Period","AccrualStart","AccrualEnd","PayDate","NotionalEUR","FixedRate(%)","DayCount","DCF","InterestEUR"]
    for c, h in enumerate(headers, 1):
        ws5.cell(3, c, h)
    style_cells(ws5, "A3:I3", header=True)
    set_col_widths(ws5, {get_column_letter(i): w for i, w in enumerate([8,12,12,12,14,12,10,10,14], 1)})
    ws5.freeze_panes = "A4"
    eur_roll: List[date] = []
    d = date(eff.year, 8, 27)
    if d <= eff:
        d = date(eff.year + 1, 2, 27)
    while d < term:
        eur_roll.append(d)
        d = add_months(d, 6)
    eur_roll.append(term)
    start = eff
    for i, end in enumerate(eur_roll, start=1):
        r = 3 + i
        pay = modified_following(end, hol_set)
        ws5[f"A{r}"], ws5[f"B{r}"], ws5[f"C{r}"], ws5[f"D{r}"] = i, start, end, pay
        ws5[f"E{r}"] = "=01_TradeTerms!C8"
        ws5[f"F{r}"] = "=01_TradeTerms!C11"
        ws5[f"G{r}"] = "ACT/360"
        ws5[f"H{r}"] = f"=(C{r}-B{r})/360"
        ws5[f"I{r}"] = f"=E{r}*(F{r}/100)*H{r}"
        for col in "BCDE":
            ws5[f"{col}{r}"].number_format = "yyyy-mm-dd"
        ws5[f"E{r}"].number_format = "#,##0"
        ws5[f"F{r}"].number_format = "0.00"
        ws5[f"H{r}"].number_format = "0.0000"
        ws5[f"I{r}"].number_format = "#,##0.00"
        style_cells(ws5, f"A{r}:I{r}")
        start = end

    # Principal exchange
    ws6 = add_sheet("06_Principal_Exchange")
    ws6["A1"] = "Principal Exchange (Final)"
    ws6["A1"].font = title_font
    ws6["A3"], ws6["B3"], ws6["C3"], ws6["D3"] = "Date","Receive HKD","Pay EUR","Locked FX (HKD/EUR)"
    style_cells(ws6, "A3:D3", header=True)
    ws6["A4"], ws6["B4"], ws6["C4"], ws6["D4"] = "=01_TradeTerms!C7", "=01_TradeTerms!C9", "=01_TradeTerms!C8", "=01_TradeTerms!C10"
    ws6["A4"].number_format = "yyyy-mm-dd"
    ws6["B4"].number_format = "#,##0"
    ws6["C4"].number_format = "#,##0"
    ws6["D4"].number_format = "0.0000"
    style_cells(ws6, "A4:D4")
    set_col_widths(ws6, {"A": 12, "B": 16, "C": 16, "D": 18})
    ws6.freeze_panes = "A4"

    # FVH effectiveness
    ws7 = add_sheet("07_Effectiveness_Template")
    ws7["A1"] = "FVH Effectiveness Template (Spot FX component)"
    ws7["A1"].font = title_font
    headers = ["ObsDate","SpotFX(HKD/EUR)","EUR_Layer","ΔSpot","ΔHI_HKD","Der_MTM_HKD(model)","Der_MTM_HKD(override)","ΔDer_HKD","X=-ΔDer","Y=ΔHI","DO_Period","CumΔHI","CumΔDer","DO_Cum","Ineff_PnL"]
    for c, h in enumerate(headers, 1):
        ws7.cell(3, c, h)
    style_cells(ws7, "A3:O3", header=True)
    set_col_widths(ws7, {get_column_letter(i): w for i, w in enumerate([12,16,12,10,14,16,18,12,12,12,10,12,12,10,12], 1)})
    ws7.freeze_panes = "A4"
    start_row = 10
    ws7["A9"], ws7["B9"] = "Initial spot", "=01_TradeTerms!C10"
    ws7["B9"].number_format = "0.0000"
    style_cells(ws7, "A9:B9")
    for i, _d in enumerate(obs):
        r = start_row + i
        ws7[f"A{r}"] = _d
        ws7[f"A{r}"].number_format = "yyyy-mm-dd"
        ws7[f"B{r}"] = f"=IF(A{r}=\"\",\"\",98_Params!$B$11*(1+98_Params!$B$12)^{i})"
        ws7[f"C{r}"] = "=02_HedgedItem!B8"
        ws7[f"F{r}"] = f"=-C{r}*B{r}"
        ws7[f"G{r}"] = ""  # override
        if i == 0:
            for col in ["D","E","H","I","J","K","L","M","N","O"]:
                ws7[f"{col}{r}"] = ""
        else:
            ws7[f"D{r}"] = f"=B{r}-B{r-1}"
            ws7[f"E{r}"] = f"=C{r}*(B{r}-B{r-1})"
            ws7[f"H{r}"] = f"=IF(G{r}<>\"\",G{r},F{r})-IF(G{r-1}<>\"\",G{r-1},F{r-1})"
            ws7[f"I{r}"] = f"=-H{r}"
            ws7[f"J{r}"] = f"=E{r}"
            ws7[f"K{r}"] = f"=IF(E{r}=0,\"\",I{r}/E{r})"
            ws7[f"L{r}"] = f"=SUM($E${start_row+1}:E{r})"
            ws7[f"M{r}"] = f"=SUM($H${start_row+1}:H{r})"
            ws7[f"N{r}"] = f"=IF(L{r}=0,\"\",-M{r}/L{r})"
            ws7[f"O{r}"] = f"=E{r}+H{r}"
        style_cells(ws7, f"A{r}:O{r}")

    # FVH regression
    ws9 = add_sheet("09_Effectiveness_Regression")
    ws9["A1"] = "FVH Regression (X=-ΔDer vs Y=ΔHI)"
    ws9["A1"].font = title_font
    ws9["A3"], ws9["B3"], ws9["C3"] = "ObsDate","X=-ΔDer_HKD","Y=ΔHI_HKD"
    style_cells(ws9, "A3:C3", header=True)
    ws9["E3"], ws9["F3"] = "Threshold","Value"
    style_cells(ws9, "E3:F3", header=True)
    ws9["E4"], ws9["F4"] = "Min R²", "=98_Params!B4"
    ws9["E5"], ws9["F5"] = "Slope lower", "=98_Params!B5"
    ws9["E6"], ws9["F6"] = "Slope upper", "=98_Params!B6"
    style_cells(ws9, "E4:F6")
    set_col_widths(ws9, {"A": 12, "B": 16, "C": 16, "E": 18, "F": 16})
    ws9.freeze_panes = "A4"
    for i in range(len(obs)):
        r = 12 + i
        src_r = start_row + i
        ws9[f"A{r}"] = f"='07_Effectiveness_Template'!A{src_r}"
        ws9[f"B{r}"] = f"='07_Effectiveness_Template'!I{src_r}"
        ws9[f"C{r}"] = f"='07_Effectiveness_Template'!E{src_r}"
        style_cells(ws9, f"A{r}:C{r}")
    end = 11 + len(obs)
    ws9["E8"], ws9["F8"] = "N", f"=COUNT(B12:B{end})"
    ws9["E9"], ws9["F9"] = "Slope", f"=IFERROR(SLOPE(C12:C{end},B12:B{end}),\"\")"
    ws9["E10"], ws9["F10"] = "Intercept", f"=IFERROR(INTERCEPT(C12:C{end},B12:B{end}),\"\")"
    ws9["E11"], ws9["F11"] = "R²", f"=IFERROR(RSQ(C12:C{end},B12:B{end}),\"\")"
    ws9["E12"], ws9["F12"] = "PASS", f"=IF(AND(F11>=F4,F9>=F5,F9<=F6),\"PASS\",\"FAIL\")"
    style_cells(ws9, "E8:F12")

    # FVH sensitivity
    ws10 = add_sheet("10_Effectiveness_Sensitivity")
    ws10["A1"] = "FVH Prospective Sensitivity (Spot shocks)"
    ws10["A1"].font = title_font
    ws10["A3"], ws10["B3"], ws10["C3"], ws10["D3"], ws10["E3"] = "Shock","Spot_Shocked","ΔHI_HKD","ΔDer_HKD","DO"
    style_cells(ws10, "A3:E3", header=True)
    shocks = [-0.10, -0.05, -0.01, 0.01, 0.05, 0.10]
    for idx, s in enumerate(shocks, start=4):
        ws10[f"A{idx}"] = s
        ws10[f"B{idx}"] = f"=01_TradeTerms!C10*(1+A{idx})"
        ws10[f"C{idx}"] = f"=02_HedgedItem!B8*(B{idx}-01_TradeTerms!C10)"
        ws10[f"D{idx}"] = f"=-C{idx}"
        ws10[f"E{idx}"] = f"=IF(C{idx}=0,\"\",-D{idx}/C{idx})"
        style_cells(ws10, f"A{idx}:E{idx}")

    # FVH summary
    ws11 = add_sheet("11_Pro_Retro_Summary")
    ws11["A1"] = "Pro/Retro Summary (FVH)"
    ws11["A1"].font = title_font
    ws11["A3"], ws11["B3"], ws11["C3"] = "Item","Value","Source"
    style_cells(ws11, "A3:C3", header=True)
    ws11["A4"], ws11["B4"] = "As-of date", f"='07_Effectiveness_Template'!A{start_row+min(10,len(obs)-1)}"
    ws11["A5"], ws11["B5"] = "Latest DO_Cum", "=LOOKUP(9E+99,'07_Effectiveness_Template'!N:N)"
    ws11["A6"], ws11["B6"] = "Regression PASS", "=09_Effectiveness_Regression!F12"
    ws11["A7"], ws11["B7"] = "Conclusion", "=IF(B6=\"PASS\",\"PASS\",\"REVIEW\")"
    for r in range(4, 8):
        style_cells(ws11, f"A{r}:C{r}")

    # CFH overview
    ws13 = add_sheet("13_CF_Hedge_Overview")
    ws13["A1"] = "Cash Flow Hedge (Optional) – Overview"
    ws13["A1"].font = title_font
    ws13["A3"], ws13["B3"] = "Item","Value"
    style_cells(ws13, "A3:B3", header=True)
    ws13["A4"], ws13["B4"] = "Designated principal for CFH (EUR)", "=02_HedgedItem!B8"
    ws13["A5"], ws13["B5"] = "Coupon rate (%)", "=02_HedgedItem!B5"
    ws13["A6"], ws13["B6"] = "Coupons per year", 2
    ws13["A7"], ws13["B7"] = "Hedge horizon end", "=01_TradeTerms!C7"
    for r in range(4, 8):
        style_cells(ws13, f"A{r}:B{r}")

    # CF interest schedule
    ws14 = add_sheet("14_CF_Interest_Schedule")
    ws14["A1"] = "Forecast EUR Interest Receipts Schedule (CFH)"
    ws14["A1"].font = title_font
    headers = ["Period","AccrualStart","CouponDate","EUR_Principal","Rate(%)","CouponEUR"]
    for c, h in enumerate(headers, 1):
        ws14.cell(3, c, h)
    style_cells(ws14, "A3:F3", header=True)
    ws14.freeze_panes = "A4"
    coupon_dates: List[date] = []
    d = date(eff.year, 8, 20)
    if d <= eff:
        d = date(eff.year + 1, 2, 20)
    while d <= term:
        coupon_dates.append(d)
        d = add_months(d, 6)

    base_row = 10
    start = eff
    for i, cd in enumerate(coupon_dates, start=1):
        r = base_row + i - 1
        ws14[f"A{r}"], ws14[f"B{r}"], ws14[f"C{r}"] = i, start, cd
        ws14[f"D{r}"] = "=13_CF_Hedge_Overview!B4"
        ws14[f"E{r}"] = "=13_CF_Hedge_Overview!B5"
        ws14[f"F{r}"] = f"=D{r}*(E{r}/100)/2"
        style_cells(ws14, f"A{r}:F{r}")
        start = cd

    # FX forwards ladder (dummy rates)
    ws15 = add_sheet("15_FX_Forwards_Ladder")
    ws15["A1"] = "FX Forwards Ladder (Template)"
    ws15["A1"].font = title_font
    headers = ["Fwd#","Maturity","HedgeEUR","FwdRate(HKD/EUR)","HedgeHKD(locked)","Notes"]
    for c, h in enumerate(headers, 1):
        ws15.cell(3, c, h)
    style_cells(ws15, "A3:F3", header=True)
    ws15.freeze_panes = "A4"
    for i in range(1, min(len(coupon_dates), 20) + 1):
        r = base_row + i - 1
        ws15[f"A{r}"] = i
        ws15[f"B{r}"] = f"=14_CF_Interest_Schedule!C{r}"
        ws15[f"C{r}"] = f"=14_CF_Interest_Schedule!F{r}"
        ws15[f"D{r}"] = f"=01_TradeTerms!C10*(1+0.001*{i})"
        ws15[f"E{r}"] = f"=C{r}*D{r}"
        ws15[f"F{r}"] = "Dummy forward rate; replace with executed"
        style_cells(ws15, f"A{r}:F{r}")

    # CF effectiveness
    ws16 = add_sheet("16_CF_Effectiveness_Template")
    ws16["A1"] = "CFH Effectiveness Template (FX forwards vs forecast coupons)"
    ws16["A1"].font = title_font
    headers = ["ObsDate","SpotFX(HKD/EUR)","ExposureEUR_Rem","ΔSpot","ΔHI_HKD","Der_MTM_HKD(model)","Der_MTM_HKD(override)","ΔDer_HKD","X=-ΔDer","Y=ΔHI","DO_Period","Effective_OCI","Ineff_PnL"]
    for c, h in enumerate(headers, 1):
        ws16.cell(3, c, h)
    style_cells(ws16, "A3:M3", header=True)
    start_row_cf = 10
    ws16["A9"], ws16["B9"] = "Initial spot", "=01_TradeTerms!C10"
    style_cells(ws16, "A9:B9")
    for i, _d in enumerate(obs):
        r = start_row_cf + i
        ws16[f"A{r}"] = _d
        ws16[f"B{r}"] = f"=IF(A{r}=\"\",\"\",98_Params!$B$11*(1+98_Params!$B$12)^{i})"
        ws16[f"C{r}"] = f"=SUMIFS('14_CF_Interest_Schedule'!$F$10:$F$199,'14_CF_Interest_Schedule'!$C$10:$C$199,\">=\"&A{r})"
        ws16[f"F{r}"] = f"=-C{r}*B{r}"
        ws16[f"G{r}"] = ""
        if i == 0:
            for col in ["D","E","H","I","J","K","L","M"]:
                ws16[f"{col}{r}"] = ""
        else:
            ws16[f"D{r}"] = f"=B{r}-B{r-1}"
            ws16[f"E{r}"] = f"=C{r}*(B{r}-B{r-1})"
            ws16[f"H{r}"] = f"=IF(G{r}<>\"\",G{r},F{r})-IF(G{r-1}<>\"\",G{r-1},F{r-1})"
            ws16[f"I{r}"] = f"=-H{r}"
            ws16[f"J{r}"] = f"=E{r}"
            ws16[f"K{r}"] = f"=IF(E{r}=0,\"\",I{r}/E{r})"
            ws16[f"L{r}"] = f"=-SIGN(E{r})*MIN(ABS(E{r}),ABS(H{r}))"
            ws16[f"M{r}"] = f"=E{r}+H{r}"
        style_cells(ws16, f"A{r}:M{r}")

    # CF regression
    ws17 = add_sheet("17_CF_Effectiveness_Regression")
    ws17["A1"] = "CFH Regression (X=-ΔDer vs Y=ΔHI)"
    ws17["A1"].font = title_font
    ws17["A3"], ws17["B3"], ws17["C3"] = "ObsDate","X=-ΔDer_HKD","Y=ΔHI_HKD"
    style_cells(ws17, "A3:C3", header=True)
    ws17["E3"], ws17["F3"] = "Threshold","Value"
    style_cells(ws17, "E3:F3", header=True)
    ws17["E4"], ws17["F4"] = "Min R²", "=98_Params!B4"
    ws17["E5"], ws17["F5"] = "Slope lower", "=98_Params!B5"
    ws17["E6"], ws17["F6"] = "Slope upper", "=98_Params!B6"
    style_cells(ws17, "E4:F6")
    for i in range(len(obs)):
        r = 12 + i
        src_r = start_row_cf + i
        ws17[f"A{r}"] = f"='16_CF_Effectiveness_Template'!A{src_r}"
        ws17[f"B{r}"] = f"='16_CF_Effectiveness_Template'!I{src_r}"
        ws17[f"C{r}"] = f"='16_CF_Effectiveness_Template'!E{src_r}"
        style_cells(ws17, f"A{r}:C{r}")
    end = 11 + len(obs)
    ws17["E8"], ws17["F8"] = "Slope", f"=IFERROR(SLOPE(C12:C{end},B12:B{end}),\"\")"
    ws17["E9"], ws17["F9"] = "R²", f"=IFERROR(RSQ(C12:C{end},B12:B{end}),\"\")"
    ws17["E10"], ws17["F10"] = "PASS", f"=IF(AND(F9>=F4,F8>=F5,F8<=F6),\"PASS\",\"FAIL\")"
    style_cells(ws17, "E8:F10")

    # CVA params + profiles + scorecard (simplified)
    ws24 = add_sheet("24_CVA_Params")
    ws24["A1"] = "CVA Parameters"
    ws24["A1"].font = title_font
    ws24["A3"], ws24["B3"], ws24["C3"] = "Parameter","Value","Notes"
    style_cells(ws24, "A3:C3", header=True)
    cva_rows = [
        ("Counterparty","=01_TradeTerms!C5",""),
        ("Rating","A","Template"),
        ("1Y PD","=98_Params!B14",""),
        ("Recovery","=98_Params!B15",""),
        ("LGD","=1-B7",""),
        ("Hazard rate λ","=-LN(1-B6)","Flat hazard"),
        ("Discount rate r","=98_Params!B16","Flat"),
        ("CSA/Collateral","No",""),
        ("As-of date","=11_Pro_Retro_Summary!B4",""),
        ("CCS maturity","=01_TradeTerms!C7",""),
    ]
    for i, (k, v, n) in enumerate(cva_rows, start=4):
        ws24[f"A{i}"], ws24[f"B{i}"], ws24[f"C{i}"] = k, v, n
        style_cells(ws24, f"A{i}:C{i}")

    def setup_cva_profile(ws, title, obs_sheet, mtm_model_col, mtm_override_col, maturity_cell, start_row_obs):
        ws["A1"] = title
        ws["A1"].font = title_font
        headers = ["ObsDate","MTM_HKD","EPE","RemYrs","RemPD","DF","CVA","ΔCVA"]
        for c, h in enumerate(headers, 1):
            ws.cell(3, c, h)
        style_cells(ws, "A3:H3", header=True)
        base = 10
        for i in range(len(obs)):
            r = base + i
            src_r = start_row_obs + i
            ws[f"A{r}"] = f"='{obs_sheet}'!A{src_r}"
            ws[f"B{r}"] = f"=IF('{obs_sheet}'!{mtm_override_col}{src_r}<>\"\",'{obs_sheet}'!{mtm_override_col}{src_r},'{obs_sheet}'!{mtm_model_col}{src_r})"
            ws[f"C{r}"] = f"=MAX(B{r},0)"
            ws[f"D{r}"] = f"=IF(A{r}=\"\",\"\",({maturity_cell}-A{r})/365)"
            ws[f"E{r}"] = f"=IF(D{r}=\"\",\"\",1-EXP(-24_CVA_Params!$B$9*D{r}))"
            ws[f"F{r}"] = f"=IF(D{r}=\"\",\"\",EXP(-24_CVA_Params!$B$10*D{r}))"
            ws[f"G{r}"] = f"=IF(A{r}=\"\",\"\",24_CVA_Params!$B$8*C{r}*E{r}*F{r})"
            ws[f"H{r}"] = f"=G{r}" if i == 0 else f"=G{r}-G{r-1}"
            style_cells(ws, f"A{r}:H{r}")
        ws["F5"], ws["G5"] = "CVA at As-of", f"=IFERROR(INDEX(G10:G{9+len(obs)},MATCH(24_CVA_Params!$B$12,A10:A{9+len(obs)},0)),\"\")"
        style_cells(ws, "F5:G5")

    ws25 = add_sheet("25_CVA_Profile_CCS")
    setup_cva_profile(ws25, "CVA Profile – CCS", "07_Effectiveness_Template", "F", "G", "24_CVA_Params!$B$13", start_row)
    ws26 = add_sheet("26_CVA_Profile_FXFWDS")
    setup_cva_profile(ws26, "CVA Profile – FX Forwards", "16_CF_Effectiveness_Template", "F", "G", "24_CVA_Params!$B$13", start_row_cf)

    ws27 = add_sheet("27_CVA_Scorecard")
    ws27["A1"] = "CVA Scorecard"
    ws27["A1"].font = title_font
    ws27["A3"], ws27["B3"], ws27["C3"] = "Metric","Value","Notes"
    style_cells(ws27, "A3:C3", header=True)
    sc = [
        ("Counterparty","=24_CVA_Params!B4",""),
        ("Rating","=24_CVA_Params!B5",""),
        ("PD(1Y)","=24_CVA_Params!B6",""),
        ("Recovery","=24_CVA_Params!B7",""),
        ("LGD","=24_CVA_Params!B8",""),
        ("CVA CCS (as-of)","=25_CVA_Profile_CCS!G5",""),
        ("CVA FXFWD (as-of)","=26_CVA_Profile_FXFWDS!G5",""),
        ("Total CVA (as-of)","=25_CVA_Profile_CCS!G5+26_CVA_Profile_FXFWDS!G5",""),
    ]
    for i, (m, v, n) in enumerate(sc, start=4):
        ws27[f"A{i}"], ws27[f"B{i}"], ws27[f"C{i}"] = m, v, n
        style_cells(ws27, f"A{i}:C{i}")

    # Journals
    headers_j = ["ObsDate","Leg","Account","SignedAmountHKD","Debit","Credit","Notes"]

    ws22 = add_sheet("22_Journal_FVH")
    ws22["A1"] = "Journal Posting – FVH (CVA in P&L)"
    ws22["A1"].font = title_font
    for c, h in enumerate(headers_j, 1):
        ws22.cell(3, c, h)
    style_cells(ws22, "A3:G3", header=True)
    out_row = 4
    for i in range(len(obs)):
        src_r = start_row + i
        prof_r = 10 + i
        obs_ref = f"'07_Effectiveness_Template'!A{src_r}"
        dDer = f"'07_Effectiveness_Template'!H{src_r}"
        dHI  = f"'07_Effectiveness_Template'!E{src_r}"
        dCVA = f"'25_CVA_Profile_CCS'!H{prof_r}"
        lines = [
            ("Derivative FV change","Derivative asset/(liab)", f"={dDer}","ΔDer"),
            ("Derivative FV change","P&L: Gain/Loss on derivative", f"=-{dDer}",""),
            ("CVA valuation adjustment","BS: CVA adjustment", f"={dCVA}","ΔCVA"),
            ("CVA valuation adjustment","P&L: CVA", f"=-{dCVA}","CVA in P&L"),
            ("Hedged item FV adjustment","Loan basis adjustment", f"={dHI}","ΔHI"),
            ("Hedged item FV adjustment","P&L: FX on hedged item", f"=-{dHI}",""),
        ]
        for leg, acct, amt, note in lines:
            ws22[f"A{out_row}"] = f"={obs_ref}"
            ws22[f"B{out_row}"], ws22[f"C{out_row}"], ws22[f"D{out_row}"], ws22[f"G{out_row}"] = leg, acct, amt, note
            ws22[f"E{out_row}"] = f"=MAX(D{out_row},0)"
            ws22[f"F{out_row}"] = f"=MAX(-D{out_row},0)"
            style_cells(ws22, f"A{out_row}:G{out_row}")
            out_row += 1

    ws23 = add_sheet("23_Journal_CFH")
    ws23["A1"] = "Journal Posting – CFH (CVA in P&L)"
    ws23["A1"].font = title_font
    for c, h in enumerate(headers_j, 1):
        ws23.cell(3, c, h)
    style_cells(ws23, "A3:G3", header=True)
    out_row = 4
    for i in range(len(obs)):
        src_r = start_row_cf + i
        prof_r = 10 + i
        obs_ref = f"'16_CF_Effectiveness_Template'!A{src_r}"
        dDer = f"'16_CF_Effectiveness_Template'!H{src_r}"
        eff  = f"'16_CF_Effectiveness_Template'!L{src_r}"
        ine  = f"'16_CF_Effectiveness_Template'!M{src_r}"
        dCVA = f"'26_CVA_Profile_FXFWDS'!H{prof_r}"
        lines = [
            ("Derivative FV change","Derivative asset/(liab)", f"={dDer}","ΔDer"),
            ("Derivative FV change","P&L: Gain/Loss on derivative", f"=-{dDer}","Gross"),
            ("CFH effective portion","OCI: CF hedge reserve", f"=-{eff}","Effective OCI"),
            ("CFH ineffectiveness","P&L: Hedge ineffectiveness", f"=-{ine}","Ineff"),
            ("CVA valuation adjustment","BS: CVA adjustment", f"={dCVA}","ΔCVA"),
            ("CVA valuation adjustment","P&L: CVA", f"=-{dCVA}","CVA in P&L"),
        ]
        for leg, acct, amt, note in lines:
            ws23[f"A{out_row}"] = f"={obs_ref}"
            ws23[f"B{out_row}"], ws23[f"C{out_row}"], ws23[f"D{out_row}"], ws23[f"G{out_row}"] = leg, acct, amt, note
            ws23[f"E{out_row}"] = f"=MAX(D{out_row},0)"
            ws23[f"F{out_row}"] = f"=MAX(-D{out_row},0)"
            style_cells(ws23, f"A{out_row}:G{out_row}")
            out_row += 1

    # Order sheets (optional)
    desired = [
        "00_Conventions","01_TradeTerms","02_HedgedItem","03_HedgeDesignation",
        "04_Schedule_HKD_Leg","05_Schedule_EUR_Leg","06_Principal_Exchange",
        "07_Effectiveness_Template","09_Effectiveness_Regression","10_Effectiveness_Sensitivity","11_Pro_Retro_Summary",
        "13_CF_Hedge_Overview","14_CF_Interest_Schedule","15_FX_Forwards_Ladder",
        "16_CF_Effectiveness_Template","17_CF_Effectiveness_Regression",
        "22_Journal_FVH","23_Journal_CFH",
        "24_CVA_Params","25_CVA_Profile_CCS","26_CVA_Profile_FXFWDS","27_CVA_Scorecard",
        "97_HK_Holidays","98_Params",
    ]
    wb._sheets = [wb[s] for s in desired if s in wb.sheetnames]

    wb.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="HKDEUR_HedgePack_from_scratch.xlsx", help="Output .xlsx path")
    args = parser.parse_args()
    out = build_hkdeur_hedgepack(args.out)
    print(f"Created: {out}")


if __name__ == "__main__":
    main()